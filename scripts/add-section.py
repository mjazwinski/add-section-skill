"""Add a new section (heading + body text) to a Word (.docx) document,
reusing the same paragraph styles (formatting) as the existing sections.

Usage:
    python add-section.py --doc <path-to-.docx> --name "Section Name" --body "Section body text"
    python add-section.py --doc <path-to-.docx> --name "Section Name" --body "Section body text" --prev-sec "Existing Section"

Notes:
 - "Body text" may contain multiple paragraphs separated by newlines ("\n" or "\r\n").
 - Each line becomes its own paragraph using the detected body style.
 - The heading style is detected by looking at the existing section headings in the
   document (paragraphs whose style name starts with "Heading"). The most commonly
   used heading level/style is reused so the new section looks consistent with the
   others. If no heading styles are found, "Heading 1" is used as a fallback.
 - The body style is detected by looking at the paragraph(s) that immediately follow
   each detected heading. The most common style among those is reused for the new
   section's body. If none is found, "Normal" is used as a fallback.
 - If --prev-sec is given, the new section is inserted immediately after the named
   section (i.e. after its last body paragraph). Raises an error if not found.
 - If --prev-sec is omitted, the new section is appended at the end of the document.
"""

import argparse
import sys
from collections import Counter

from docx import Document


def detect_heading_style(document):
    """Return the most frequently used heading style name among existing sections."""
    heading_styles = [
        p.style.name
        for p in document.paragraphs
        if p.style and p.style.name and p.style.name.lower().startswith("heading") and p.text.strip()
    ]
    if not heading_styles:
        return "Heading 1"
    return Counter(heading_styles).most_common(1)[0][0]


def detect_body_style(document, heading_style_name):
    """Return the most common style of the first non-empty paragraph after each
    heading that uses heading_style_name."""
    paragraphs = document.paragraphs
    body_styles = []
    for idx, p in enumerate(paragraphs):
        if p.style.name == heading_style_name and p.text.strip():
            for follower in paragraphs[idx + 1:]:
                if follower.style and follower.style.name.lower().startswith("heading"):
                    break
                if follower.text.strip():
                    body_styles.append(follower.style.name)
                    break
    if not body_styles:
        return "Normal"
    return Counter(body_styles).most_common(1)[0][0]


def _find_section_tail(document, section_name, heading_style_name):
    """Return the last paragraph of the named section.

    Searches for a heading paragraph whose text matches section_name (case-insensitive).
    Returns the last paragraph before the next heading (or end of document).
    Returns None if the section heading is not found.
    """
    paragraphs = document.paragraphs
    found_idx = None
    for idx, p in enumerate(paragraphs):
        if (
            p.style
            and p.style.name == heading_style_name
            and p.text.strip().lower() == section_name.strip().lower()
        ):
            found_idx = idx
            break
    if found_idx is None:
        return None
    tail = paragraphs[found_idx]
    for p in paragraphs[found_idx + 1:]:
        if p.style and p.style.name and p.style.name.lower().startswith("heading"):
            break
        tail = p
    return tail


def add_section(doc_path, section_name, section_body, heading_style=None, prev_section=None):
    document = Document(doc_path)

    heading_style = heading_style or detect_heading_style(document)
    body_style = detect_body_style(document, heading_style)

    lines = section_body.splitlines() or [section_body]

    if prev_section is not None:
        ref_para = _find_section_tail(document, prev_section, heading_style)
        if ref_para is None:
            raise ValueError(
                f"Section '{prev_section}' not found in document "
                f"(looked for a '{heading_style}' paragraph with that text)."
            )
        # Add heading at document end then move it right after ref_para.
        new_heading = document.add_paragraph(section_name, style=heading_style)
        ref_para._p.addnext(new_heading._p)
        # Insert each body line after the previous element.
        prev_elem = new_heading._p
        for line in lines:
            body_para = document.add_paragraph(line, style=body_style)
            prev_elem.addnext(body_para._p)
            prev_elem = body_para._p
    else:
        document.add_paragraph(section_name, style=heading_style)
        for line in lines:
            document.add_paragraph(line, style=body_style)

    document.save(doc_path)
    location = f"after section '{prev_section}'" if prev_section else "at end of document"
    print(
        f"Added section '{section_name}' {location} "
        f"using heading style '{heading_style}' and body style '{body_style}'."
    )


def main():
    parser = argparse.ArgumentParser(description="Add a new section to a Word document.")
    parser.add_argument("--doc", required=True, help="Path to the .docx document to edit")
    parser.add_argument("--name", required=True, help="Name/title of the new section")
    parser.add_argument("--body", required=True, help="Body text of the new section")
    parser.add_argument(
        "--heading-style",
        default=None,
        help="Force a specific heading style (e.g. 'Heading 1') instead of "
             "auto-detecting the most common one used in the document",
    )
    parser.add_argument(
        "--prev-sec",
        default=None,
        help="Name of the existing section after which the new section should be inserted. "
             "If omitted, the new section is appended at the end of the document.",
    )
    args = parser.parse_args()

    try:
        add_section(args.doc, args.name, args.body, args.heading_style, args.prev_sec)
    except Exception as exc:  # surface a clear error to the caller
        print(f"Error adding section: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
